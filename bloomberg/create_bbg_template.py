"""Generate a Bloomberg Excel template for SG historical pulls.

Output matches the four raw CSVs expected by this project:
  price_data.csv, book_data.csv, dividend_data.csv, categorizations_data.csv

Requires Bloomberg Terminal + Excel Add-in to refresh formulas.
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

from config import SG_TICKERS, MARKETS


OUT_PATH = Path(__file__).resolve().parent / "SG_Bloomberg_Historical_Pull.xlsx"
OUT_PATH_FALLBACK = Path(__file__).resolve().parent / "SG_Bloomberg_Historical_Pull_5y.xlsx"
START_DATE = MARKETS["sg"].start_date  # YYYY-MM-DD
BENCHMARK_YAHOO = MARKETS["sg"].benchmark_symbol  # ES3.SI


def yahoo_to_bbg(yahoo: str) -> str:
    """D05.SI → D05 SP Equity; ES3.SI → ES3 SP Equity."""
    root = yahoo.replace(".SI", "").strip().upper()
    return f"{root} SP Equity"


def style_header(cell, fill_hex: str = "1F4E79") -> None:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=fill_hex)
    cell.alignment = Alignment(horizontal="center", wrap_text=True)


def thin_border():
    side = Side(style="thin", color="B0B0B0")
    return Border(left=side, right=side, top=side, bottom=side)


def build_workbook() -> Workbook:
    wb = Workbook()

    # ── 0) README ──────────────────────────────────────────────
    ws = wb.active
    ws.title = "README"
    lines = [
        ("SG Bloomberg Historical Pull — Quant Trading", True),
        ("", False),
        ("Purpose", True),
        (
            "Pull SGX history from Bloomberg into the 4 CSV layouts this project needs "
            "under data/sg/raw/, then run: python main.py --market sg --skip-download",
            False,
        ),
        ("", False),
        ("Sheets", True),
        ("1. Control — start/end dates and pull settings (edit these first)", False),
        ("2. Universe — Yahoo .SI tickers ↔ Bloomberg SP Equity codes", False),
        ("3. Prices — each ticker Date|Value BDH block (→ price_data.csv)", False),
        ("4. Book — each ticker Date|Value BDH Per=M (→ book_data.csv)", False),
        ("5. Categories — sector + market cap BDP (→ categorizations_data.csv)", False),
        ("6. Dividends — BDS DVD_HIST_ALL for EVERY ticker (side-by-side blocks)", False),
        ("7. Dividends_Detail — one-ticker DVD_HIST + Regular vs Special sums", False),
        ("8. FAQ — BQ error / Special dividends / why DBS", False),
        ("9. Export_Checklist — save-as-CSV steps", False),
        ("", False),
        ("How to refresh (Bloomberg Terminal PC)", True),
        ("1. Open this file in Excel with Bloomberg Add-in logged in", False),
        ("2. Edit Control!B2 (Start) and Control!B3 (End)", False),
        ("3. Press Bloomberg → Refresh Workbook (or wait for auto-refresh)", False),
        ("4. Follow Export_Checklist, then run bloomberg/bbg_to_csv.py", False),
        ("", False),
        ("Field mapping", True),
        ("Prices: PX_LAST with CshAdjNormal=Y, CshAdjAbnormal=Y, CapChg=Y (≈ Yahoo auto_adjust)", False),
        ("Book: BOOK_VAL_PER_SH with Per=M (monthly; NOT BQ)", False),
        ("Categories: GICS_SECTOR_NAME + CUR_MKT_CAP", False),
        ("Dividends: BDS DVD_HIST_ALL (Bloomberg dividend history) for all tickers", False),
        ("", False),
        ("Notes", True),
        ("Column headers stay as Yahoo .SI names so CSVs drop straight into the pipeline.", False),
        ("Benchmark ES3.SI is on Universe for reference; strategy downloads it separately.", False),
        ("TA-Lib / OHLCV not required — this project uses adjusted close only.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        ws.cell(i, 1, text).font = Font(bold=bold, size=12 if bold and i == 1 else 11)
    ws.column_dimensions["A"].width = 110

    # ── 1) Control ─────────────────────────────────────────────
    ctrl = wb.create_sheet("Control")
    ctrl["A1"] = "Parameter"
    ctrl["B1"] = "Value"
    style_header(ctrl["A1"])
    style_header(ctrl["B1"])

    params = [
        (
            "StartDate",
            "=EDATE(TODAY(),-60)",
            "Rolling 5 years (60 months). Hardcode e.g. 2021-09-01 to match config if preferred.",
        ),
        ("EndDate", "=TODAY()-1", "Leave as TODAY()-1 or hardcode e.g. 2026-09-04"),
        ("PriceField", "PX_LAST", "Adjusted via BDH overrides on Prices sheet"),
        ("BookField", "BOOK_VAL_PER_SH", "Book value per share (BDH)"),
        ("BookPeriod", "M", "Monthly BDH periodicity for book value (was Q)"),
        ("PricePeriod", "D", "Daily"),
        ("DivField", "DVD_HIST_ALL", "Bloomberg dividend history via BDS (not BDH)"),
        ("Currency", "SGD", "Informational"),
        ("Exchange", "SP / SGX", "Bloomberg yellow-key Equity, SGX"),
        ("OutputFolder", str(ROOT / "data" / "sg" / "raw"), "Where CSVs should land"),
    ]
    for r, (name, val, note) in enumerate(params, start=2):
        ctrl.cell(r, 1, name)
        cell = ctrl.cell(r, 2, val)
        if isinstance(val, str) and val.startswith("="):
            cell.value = val
        ctrl.cell(r, 3, note)

    ctrl["A12"] = "BDH price overrides (already embedded in Prices formulas)"
    ctrl["A13"] = "CshAdjNormal=Y | CshAdjAbnormal=Y | CapChg=Y | Dir=V | Dts=S | Sort=A | Fill=P | Days=A"
    ctrl.column_dimensions["A"].width = 18
    ctrl.column_dimensions["B"].width = 22
    ctrl.column_dimensions["C"].width = 70

    # Named-style refs used in formulas: Control!$B$2, Control!$B$3
    # EndDate formula stays live.

    # ── 2) Universe ────────────────────────────────────────────
    uni = wb.create_sheet("Universe")
    headers = [
        "YahooTicker",
        "BloombergTicker",
        "Include",
        "Name_BDP",
        "Last_BDP",
        "Sector_BDP",
        "MktCap_BDP",
    ]
    for c, h in enumerate(headers, start=1):
        style_header(uni.cell(1, c, h))

    tickers = list(SG_TICKERS)
    # Ensure benchmark listed (for reference / optional price column)
    if BENCHMARK_YAHOO not in tickers:
        tickers_with_bench = tickers + [BENCHMARK_YAHOO]
    else:
        tickers_with_bench = tickers

    for i, yahoo in enumerate(tickers_with_bench, start=2):
        bbg = yahoo_to_bbg(yahoo)
        uni.cell(i, 1, yahoo)
        uni.cell(i, 2, bbg)
        # Include = YES for strategy universe; NO for benchmark-only rows if you prefer
        uni.cell(i, 3, "YES" if yahoo in SG_TICKERS else "BENCH")
        # Live BDP formulas
        uni.cell(i, 4, f'=BDP(B{i},"NAME")')
        uni.cell(i, 5, f'=BDP(B{i},"PX_LAST")')
        uni.cell(i, 6, f'=BDP(B{i},"GICS_SECTOR_NAME")')
        uni.cell(i, 7, f'=BDP(B{i},"CUR_MKT_CAP")')

    for col, width in zip("ABCDEFG", [12, 16, 10, 28, 12, 22, 16]):
        uni.column_dimensions[col].width = width

    n_uni = len(tickers_with_bench)

    # ── 3) Prices — each ticker has its OWN Date|Value BDH (Dts=S) ─
    # Old layout only showed dates for DBS (first ticker). Now every name
    # gets a 2-column block: Date + PX_LAST, same pattern as Dividends.
    COLS_PER_PX = 2
    prices = wb.create_sheet("Prices")
    prices["A1"] = (
        "Each ticker = 2 columns (Date | Value) with Dts=S. NOT DBS-only dates. "
        "Scroll right. Run bbg_to_csv.py to merge into wide price_data.csv."
    )
    prices["A1"].font = Font(bold=True, size=11)
    prices.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)

    px_idx = wb.create_sheet("Prices_Index")
    for c, h in enumerate(
        ["YahooTicker", "BloombergTicker", "StartCol", "ExcelCol", "BDH_Cell"],
        start=1,
    ):
        style_header(px_idx.cell(1, c, h), "0D7377")

    for i, yahoo in enumerate(tickers):
        uni_row = tickers_with_bench.index(yahoo) + 2
        start_col = 1 + i * COLS_PER_PX
        label_row, formula_row = 3, 4
        style_header(prices.cell(label_row, start_col, yahoo), "0D7377")
        prices.cell(label_row, start_col + 1, "Value")
        style_header(prices.cell(label_row, start_col + 1), "0D7377")
        prices.cell(
            formula_row,
            start_col,
            (
                f'=BDH(Universe!B{uni_row},"PX_LAST",Control!$B$2,Control!$B$3,'
                f'"Dir=V","Dts=S","Sort=A","CshAdjNormal=Y","CshAdjAbnormal=Y",'
                f'"CapChg=Y","Fill=P","Days=A")'
            ),
        )
        r = i + 2
        px_idx.cell(r, 1, yahoo)
        px_idx.cell(r, 2, yahoo_to_bbg(yahoo))
        px_idx.cell(r, 3, start_col)
        px_idx.cell(r, 4, get_column_letter(start_col))
        px_idx.cell(r, 5, f"{get_column_letter(start_col)}{formula_row}")

    for col, w in zip("ABCDE", [12, 16, 10, 10, 12]):
        px_idx.column_dimensions[col].width = w

    note = wb.create_sheet("Prices_Notes")
    note["A1"] = "Prices sheet notes"
    note["A1"].font = Font(bold=True, size=12)
    note_lines = [
        "Each ticker has its OWN dates (Dts=S) — not only DBS.",
        "Block layout: Row3 = Yahoo ticker | Value; Row4 = BDH spills Date + PX_LAST.",
        "See Prices_Index for column map (D05 starts col A, O39 starts col C, …).",
        "bbg_to_csv.py outer-joins all blocks on Date → wide price_data.csv.",
        "Fill=P / Days=A: edit formula if you want exchange trading days only.",
    ]
    for i, t in enumerate(note_lines, start=2):
        note.cell(i, 1, t)
    note.column_dimensions["A"].width = 100

    # ── 4) Book — each ticker OWN Date|Value BDH, Per=M ─────────
    COLS_PER_BOOK = 2
    book = wb.create_sheet("Book")
    book["A1"] = (
        "Each ticker = 2 columns (Date | BOOK_VAL_PER_SH), Per=M, Dts=S. "
        "Every name has its own dates (not DBS-only)."
    )
    book["A1"].font = Font(bold=True, size=11)
    book.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)

    book_idx = wb.create_sheet("Book_Index")
    for c, h in enumerate(
        ["YahooTicker", "BloombergTicker", "StartCol", "ExcelCol", "BDH_Cell"],
        start=1,
    ):
        style_header(book_idx.cell(1, c, h), "6B3FA0")

    for i, yahoo in enumerate(tickers):
        uni_row = tickers_with_bench.index(yahoo) + 2
        start_col = 1 + i * COLS_PER_BOOK
        label_row, formula_row = 3, 4
        style_header(book.cell(label_row, start_col, yahoo), "6B3FA0")
        style_header(book.cell(label_row, start_col + 1, "Value"), "6B3FA0")
        book.cell(
            formula_row,
            start_col,
            (
                f'=BDH(Universe!B{uni_row},Control!$B$5,Control!$B$2,Control!$B$3,'
                f'"Dir=V","Dts=S","Sort=A","Per=M","Fill=P")'
            ),
        )
        r = i + 2
        book_idx.cell(r, 1, yahoo)
        book_idx.cell(r, 2, yahoo_to_bbg(yahoo))
        book_idx.cell(r, 3, start_col)
        book_idx.cell(r, 4, get_column_letter(start_col))
        book_idx.cell(r, 5, f"{get_column_letter(start_col)}{formula_row}")

    for col, w in zip("ABCDE", [12, 16, 10, 10, 12]):
        book_idx.column_dimensions[col].width = w

    book_note = wb.create_sheet("Book_Notes")
    book_note["A1"] = "Book sheet notes"
    book_note["A1"].font = Font(bold=True)
    for i, t in enumerate(
        [
            "WHY YOU SAW 'Enumeration Error' WITH BQ:",
            "  Bloomberg BDH periodicity must be D / W / M / Q / Y (or CQ / FY etc.).",
            "  'BQ' is NOT valid → Enumeration Error. This template now uses Per=M.",
            "",
            "DATES: every ticker has its own Date column (Dts=S), not only DBS.",
            "Field: BOOK_VAL_PER_SH, Period: M (monthly).",
            "See Book_Index for column map. bbg_to_csv.py merges blocks → book_data.csv.",
            "Fundamentals may only update quarterly; monthly series can repeat last report.",
        ],
        start=2,
    ):
        book_note.cell(i, 1, t)
    book_note.column_dimensions["A"].width = 100

    # ── 5) Categories (static BDP snapshot) ─────────────────────
    cat = wb.create_sheet("Categories")
    for c, h in enumerate(["Ticker", "sector", "market_cap"], start=1):
        style_header(cat.cell(1, c, h), "C45C26")

    for i, yahoo in enumerate(tickers, start=2):
        uni_row = tickers_with_bench.index(yahoo) + 2
        cat.cell(i, 1, yahoo)  # keep Yahoo name as Ticker index
        cat.cell(i, 2, f'=BDP(Universe!B{uni_row},"GICS_SECTOR_NAME")')
        cat.cell(i, 3, f'=BDP(Universe!B{uni_row},"CUR_MKT_CAP")')

    for col, w in zip("ABC", [12, 24, 16]):
        cat.column_dimensions[col].width = w

    cat_note = wb.create_sheet("Categories_Notes")
    cat_note["A1"] = "Categories notes"
    cat_note["A1"].font = Font(bold=True)
    cat_note["A2"] = "Export this sheet as categorizations_data.csv (headers: Ticker,sector,market_cap)."
    cat_note["A3"] = "If GICS_SECTOR_NAME is empty, switch formula to INDUSTRY_SECTOR or CLASSIFICATION_LEVEL_1_NAME."
    cat_note["A4"] = "market_cap should be absolute CUR_MKT_CAP (same units across names)."
    cat_note.column_dimensions["A"].width = 100

    # ── 6) Dividends — BDS DVD_HIST_ALL for ALL tickers ─────────
    # DVD_HIST / DVD_HIST_ALL is a BDS field (not BDH). Each ticker gets its own
    # horizontal block so spills do not overwrite each other.
    COLS_PER_DIV = 8  # room for Ex-Date, Amount, Type, etc.

    div = wb.create_sheet("Dividends")
    div["A1"] = (
        "DVD_HIST_ALL (BDS) — one block per ticker. Row3=Yahoo ticker, Row4=BDS formula. "
        "Refresh Bloomberg, then run bbg_to_csv.py"
    )
    div["A1"].font = Font(bold=True, size=11)
    div.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)

    # Index sheet: where each ticker block starts
    div_idx = wb.create_sheet("Dividends_Index")
    for c, h in enumerate(
        ["YahooTicker", "BloombergTicker", "StartCol", "ExcelCol", "BDS_Cell"],
        start=1,
    ):
        style_header(div_idx.cell(1, c, h), "2E7D32")

    for i, yahoo in enumerate(tickers):
        uni_row = tickers_with_bench.index(yahoo) + 2
        start_col = 1 + i * COLS_PER_DIV
        # Leave row 1 as banner; ticker blocks start at sheet row 3
        label_row, formula_row = 3, 4
        label_cell = div.cell(label_row, start_col, yahoo)
        style_header(label_cell, "2E7D32")
        # BDS DVD_HIST_ALL — use DVD_START_DT / DVD_END_DT (NOT StartDateTime/EndDateTime)
        div.cell(
            formula_row,
            start_col,
            (
                f'=BDS(Universe!B{uni_row},Control!$B$8,"Headers=Y",'
                f'"DVD_START_DT="&TEXT(Control!$B$2,"yyyymmdd"),'
                f'"DVD_END_DT="&TEXT(Control!$B$3,"yyyymmdd"))'
            ),
        )
        # Index
        r = i + 2
        div_idx.cell(r, 1, yahoo)
        div_idx.cell(r, 2, yahoo_to_bbg(yahoo))
        div_idx.cell(r, 3, start_col)
        div_idx.cell(r, 4, get_column_letter(start_col))
        div_idx.cell(r, 5, f"{get_column_letter(start_col)}{formula_row}")

    for col, w in zip("ABCDE", [12, 16, 10, 10, 12]):
        div_idx.column_dimensions[col].width = w

    div_note = wb.create_sheet("Dividends_Notes")
    div_note["A1"] = "Dividends use Bloomberg DVD_HIST (BDS DVD_HIST_ALL)"
    div_note["A1"].font = Font(bold=True, size=12)
    for i, t in enumerate(
        [
            "",
            "Function: BDS(security, \"DVD_HIST_ALL\", \"Headers=Y\", DVD_START_DT, DVD_END_DT)",
            "Date overrides MUST be DVD_START_DT / DVD_END_DT (yyyymmdd).",
            "Do NOT use StartDateTime / EndDateTime — those cause: Invalid override field: EndDateTime.",
            "Control!B8 = DVD_HIST_ALL (or DVD_HIST if your terminal prefers that field).",
            "",
            "Layout:",
            "  • Dividends sheet: each ticker owns 8 columns (side-by-side) so histories do not clash.",
            "  • Row 3 = Yahoo ticker label (D05.SI, O39.SI, …) — NOT only DBS.",
            "  • Row 4 = BDS formula; spill shows Ex-Date, Dividend Amount, Dividend Type, …",
            "  • Dividends_Index lists every ticker and its starting column.",
            "",
            "Special dividends:",
            "  DVD_HIST lists Type = Regular / Interim / Final / Special / Stock separately.",
            "  It does not auto-sum. bbg_to_csv.py sums cash DPS per Ex-Date and ticker,",
            "  drops Stock dividends, keeps Special cash (Yahoo-like). Edit converter to exclude Special if needed.",
            "",
            "After refresh: python bloomberg/bbg_to_csv.py --xlsx bloomberg/SG_Bloomberg_Historical_Pull.xlsx",
        ],
        start=2,
    ):
        div_note.cell(i, 1, t)
    div_note.column_dimensions["A"].width = 110

    # Optional detail helper (single ticker) with SUMIF for Special vs Regular
    detail = wb.create_sheet("Dividends_Detail")
    detail["A1"] = "Optional: one-ticker DVD_HIST_ALL — inspect Special vs Regular"
    detail["A1"].font = Font(bold=True, size=12)
    detail["A3"] = "SelectYahoo"
    detail["B3"] = tickers[0]
    detail["C3"] = "← change me (e.g. O39.SI). Default D05.SI is DBS only as an example."
    detail["A4"] = "BloombergTicker"
    detail["B4"] = '=IFERROR(INDEX(Universe!B:B,MATCH(B3,Universe!A:A,0)),"")'
    detail["A6"] = "BDS DVD_HIST_ALL (spills below)"
    detail["A7"] = (
        '=BDS(B4,Control!$B$8,"Headers=Y",'
        '"DVD_START_DT="&TEXT(Control!$B$2,"yyyymmdd"),'
        '"DVD_END_DT="&TEXT(Control!$B$3,"yyyymmdd"))'
    )
    detail["A9"] = "After refresh, point these SUMIFs at your Amount and Type columns in the spill:"
    detail["A10"] = "Sum ALL amounts (example B8:B500):"
    detail["B10"] = "=IFERROR(SUM(B8:B500),0)"
    detail["A11"] = "Sum Ordinary (Regular/Interim/Final):"
    detail["B11"] = (
        '=IFERROR(SUMIF(C8:C500,"*Regular*",B8:B500)'
        '+SUMIF(C8:C500,"*Interim*",B8:B500)'
        '+SUMIF(C8:C500,"*Final*",B8:B500),0)'
    )
    detail["A12"] = "Sum Special only:"
    detail["B12"] = '=IFERROR(SUMIF(C8:C500,"*Special*",B8:B500),0)'
    detail["A14"] = "Adjust column letters to match your BDS spill (Amount / Type)."
    for col, w in zip("ABC", [70, 18, 70]):
        detail.column_dimensions[col].width = w

    # ── 7) FAQ ─────────────────────────────────────────────────
    faq = wb.create_sheet("FAQ")
    faq["A1"] = "FAQ — Book BQ error / Dividends Special / Only DBS dates"
    faq["A1"].font = Font(bold=True, size=12)
    faq_lines = [
        "",
        "1) Book value Enumeration Error and BQ",
        "   Cause: Per=BQ is invalid in Bloomberg BDH.",
        "   Fix: use Per=M (monthly) or Per=Q (quarterly). Never BQ.",
        "",
        "2) Dividend history — use DVD_HIST (BDS DVD_HIST_ALL)",
        "   Overrides: DVD_START_DT / DVD_END_DT (yyyymmdd). Never EndDateTime.",
        "   Dividends sheet has one BDS block per ticker (see Dividends_Index).",
        "   Special = dividend type line, not a sum. Use Dividends_Detail SUMIF for totals.",
        "",
        "3) Why only DBS dates before?",
        "   Old Prices/Book used Dts=S only on first ticker (D05.SI / DBS).",
        "   Now EVERY ticker has its own Date|Value block (Dts=S). Scroll right / see Index sheets.",
    ]
    for i, t in enumerate(faq_lines, start=2):
        faq.cell(i, 1, t)
    faq.column_dimensions["A"].width = 100

    # ── 8) Export checklist ────────────────────────────────────
    exp = wb.create_sheet("Export_Checklist")
    exp["A1"] = "Export checklist → data/sg/raw/"
    exp["A1"].font = Font(bold=True, size=12)
    steps = [
        "",
        "After all sheets show numbers (not #N/A Requesting Data...):",
        "",
        "Option A — Automated (recommended)",
        "  1. Save this workbook (formulas calculated / values cached).",
        "  2. Confirm Prices/Book: each ticker has its own Date column (scroll right).",
        "  3. Refresh Dividends (DVD_HIST_ALL blocks).",
        "  4. Run:",
        "       python bloomberg/bbg_to_csv.py --xlsx bloomberg/SG_Bloomberg_Historical_Pull.xlsx",
        "  5. Confirm data/sg/raw/*.csv then: python main.py --market sg --skip-download",
        "",
        "Validation",
        "  • Prices/Book: per-ticker Date|Value blocks (not DBS-only dates)",
        "  • Dividends: BDS DVD_HIST_ALL + DVD_START_DT / DVD_END_DT",
        "  • Book: Per=M",
    ]
    for i, t in enumerate(steps, start=2):
        exp.cell(i, 1, t)
    exp.column_dimensions["A"].width = 100

    return wb


def main() -> None:
    wb = build_workbook()
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(OUT_PATH)
        saved = OUT_PATH
    except PermissionError:
        wb.save(OUT_PATH_FALLBACK)
        saved = OUT_PATH_FALLBACK
        print(f"Note: {OUT_PATH.name} is open — saved as {OUT_PATH_FALLBACK.name} instead.")
    print(f"Wrote {saved}")
    print(f"Tickers: {len(SG_TICKERS)} | Config start: {START_DATE} | Excel StartDate: EDATE(TODAY(),-60) = 5y")
    print("Open in Excel on a Bloomberg Terminal PC and refresh.")


if __name__ == "__main__":
    main()
